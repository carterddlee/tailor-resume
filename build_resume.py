"""Render tailored resume JSON files into .docx.

Usage:
  python build_resume.py                    # builds every output/*/resume.json
  python build_resume.py path/to/resume.json [...]
"""
import glob
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Calibri"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
RIGHT_TAB = Inches(7.5)  # page width 8.5" minus 0.5" margins


def setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(0.5))
    sec.top_margin = sec.bottom_margin = Inches(0.5)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)


def para(doc, space_before=0, space_after=0, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in {"w:val": "single", "w:sz": "6", "w:space": "1", "w:color": "1F3A5F"}.items():
        b.set(qn(k), v)
    bdr.append(b)
    pPr.append(bdr)


def section_header(doc, text):
    p = para(doc, space_before=8, space_after=3)
    r = p.add_run(text.upper())
    r.bold, r.font.size, r.font.color.rgb = True, Pt(11), ACCENT
    bottom_border(p)


def left_right(doc, left, right, bold_left=True, italic=False, space_before=0):
    p = para(doc, space_before=space_before)
    p.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run(left)
    r.bold, r.italic = bold_left, italic
    if right:
        r2 = p.add_run("\t" + right)
        r2.italic = italic
    return p


def bullets(doc, items):
    for text in items:
        p = para(doc, style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.25)
        p.add_run(text)


def build(data, out_path):
    doc = Document()
    setup(doc)

    # Header
    p = para(doc)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(data["name"])
    r.bold, r.font.size, r.font.color.rgb = True, Pt(18), ACCENT
    p = para(doc, space_after=4)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("  |  ".join(data.get("contact", []))).font.size = Pt(9.5)

    if data.get("summary"):
        section_header(doc, "Summary")
        para(doc).add_run(data["summary"])

    if data.get("experience"):
        section_header(doc, "Experience")
        for i, job in enumerate(data["experience"]):
            left_right(doc, job["company"], job.get("location", ""), space_before=0 if i == 0 else 4)
            left_right(doc, job["title"], job.get("dates", ""), bold_left=False, italic=True)
            bullets(doc, job.get("bullets", []))

    if data.get("projects"):
        section_header(doc, "Projects")
        for i, proj in enumerate(data["projects"]):
            p = para(doc, space_before=0 if i == 0 else 4)
            p.add_run(proj["name"]).bold = True
            if proj.get("tech"):
                p.add_run(f"  |  {proj['tech']}").italic = True
            bullets(doc, proj.get("bullets", []))

    if data.get("education"):
        section_header(doc, "Education")
        for ed in data["education"]:
            left_right(doc, ed["school"], ed.get("dates", ""))
            para(doc).add_run(ed["degree"] + (f" — {ed['details']}" if ed.get("details") else ""))

    if data.get("skills"):
        section_header(doc, "Skills")
        for group, items in data["skills"].items():
            p = para(doc)
            p.add_run(f"{group}: ").bold = True
            p.add_run(", ".join(items))

    if data.get("certifications"):
        section_header(doc, "Certifications")
        for c in data["certifications"]:
            para(doc).add_run(c)

    doc.save(out_path)


def main():
    paths = sys.argv[1:] or glob.glob("output/*/resume.json")
    if not paths:
        print("No resume.json files found in output/.")
        return
    for path in paths:
        path = Path(path)
        try:
            data = json.loads(path.read_text())
            out = path.parent / f"{path.parent.name}-resume.docx"
            build(data, out)
            print(f"OK   {out}")
        except Exception as e:
            print(f"FAIL {path}: {e}")


if __name__ == "__main__":
    main()
