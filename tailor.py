"""Format-aware helpers for /tailor. The base resume can be any file type;
each tailored resume comes back in the same type and keeps the base's formatting.

Usage:
  python3 tailor.py prepare            # detect base/<file>, extract editable content
  python3 tailor.py build [slug ...]   # write output/<slug>/carter_lee_<company>_resume.<ext>

Two modes:
  text   (.md .txt .tex .html .json .yaml ...): the agent edits a copy of the file
         directly and writes output/<slug>/tailored.<ext>; build renames it.
  blocks (.docx .pdf .doc .rtf .odt .pages): prepare turns the base into
         base/.work/base.docx and lists its paragraphs in base/.work/content.json.
         The agent writes output/<slug>/edits.json, and build applies those edits to
         a copy of base.docx (so fonts, spacing, bullets and layout carry over)
         and converts the result back to the original file type.
  Either way the agent also writes output/<slug>/job.json {"company", "role"},
  which build uses to name the file.

edits.json:
  {
    "replace": {"12": "new text", "15": ["Bold lead-in: ", "rest of bullet"]},
    "reorder": [[20, 18, 19]],   # these paragraphs swap into each other's slots
    "delete":  [21, 22]
  }
  A list value in "replace" gives one string per formatting segment ("runs" in
  content.json), so mixed bold/italic inside a paragraph is preserved.
"""
import copy
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from lxml import etree
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
BASE_DIR = ROOT / "base"
WORK = BASE_DIR / ".work"
OUTPUT = ROOT / "output"
NAME_PREFIX = "carter_lee"  # tailored files are named carter_lee_<company>_resume.<ext>

TEXT_EXTS = {".md", ".markdown", ".txt", ".tex", ".html", ".htm", ".json", ".yaml",
             ".yml", ".rst", ".org", ".typ", ".adoc", ".xml", ".csv"}
TEXTUTIL_EXTS = {".doc", ".rtf", ".odt"}  # round-trip through macOS textutil
BLOCK_EXTS = {".docx", ".pdf", ".pages"} | TEXTUTIL_EXTS

W_P, W_R, W_T = qn("w:p"), qn("w:r"), qn("w:t")
TEXT_TAGS = {qn("w:t"): None, qn("w:tab"): "\t", qn("w:br"): "\n", qn("w:cr"): "\n",
             qn("w:noBreakHyphen"): "-", qn("w:softHyphen"): ""}
MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"


# ---------- locating the base resume ----------

def find_base():
    files = sorted(p for p in BASE_DIR.iterdir()
                   if p.is_file() and not p.name.startswith((".", "~$"))
                   and "-preview." not in p.name)
    if not files:
        sys.exit("No base resume found. Put your resume (any file type) in base/.")
    if len(files) > 1:
        sys.exit("Found more than one file in base/; keep only one:\n  "
                 + "\n  ".join(str(f.relative_to(ROOT)) for f in files))
    base = files[0]
    ext = base.suffix.lower()
    if ext not in TEXT_EXTS | BLOCK_EXTS:
        sys.exit(f"Unsupported base resume type '{ext}'. Supported: "
                 + " ".join(sorted(TEXT_EXTS | BLOCK_EXTS)))
    return base, ext


# ---------- conversions ----------

def run(cmd, timeout=180):
    subprocess.run(cmd, check=True, timeout=timeout, capture_output=True)


def soffice():
    for c in ("soffice", "libreoffice",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if shutil.which(c) or Path(c).exists():
            return shutil.which(c) or c
    return None


WORD_APP = Path("/Applications/Microsoft Word.app")
WORD_FORMATS = {".docx": "format document", ".doc": "format document97",
                ".rtf": "format rtf", ".pdf": "format PDF"}


def word_save_as(src, dst):
    """Open src in Microsoft Word and save it as dst's file type. Word is sandboxed,
    so work inside its container folder to avoid file-access prompts."""
    box = Path.home() / "Library/Containers/com.microsoft.Word/Data/tailor-resume"
    box.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=box) as tmp:
        tmp_src = Path(tmp) / ("in" + src.suffix)
        tmp_dst = Path(tmp) / ("out" + dst.suffix)
        shutil.copy(src, tmp_src)
        osascript(f'''on run argv
          tell application "Microsoft Word"
            open file name (item 1 of argv)
            save as active document file name (item 2 of argv) file format {WORD_FORMATS[dst.suffix.lower()]}
            close active document saving no
          end tell
        end run''', tmp_src, tmp_dst)
        shutil.move(str(tmp_dst), str(dst))


def to_docx(src, dst):
    ext = src.suffix.lower()
    if ext == ".docx":
        shutil.copy(src, dst)
    elif ext == ".pdf":
        from pdf2docx import Converter  # pip install pdf2docx
        cv = Converter(str(src))
        try:
            cv.convert(str(dst))
        finally:
            cv.close()
        restore_pdf_fonts(src, dst)
    elif ext in WORD_FORMATS and WORD_APP.exists():
        word_save_as(src, dst)
    elif ext in TEXTUTIL_EXTS:
        run(["textutil", "-convert", "docx", str(src), "-output", str(dst)])
    elif ext == ".pages":
        pages_export(src, dst)
    else:
        raise ValueError(ext)


def font_family(pdf_font):
    """'ABCDEF+TimesNewRomanPS-BoldMT' -> 'Times New Roman'"""
    name = pdf_font.split("+")[-1].split("-")[0].split(",")[0]
    name = re.sub(r"(PS)?MT$|PS$", "", name)
    return re.sub(r"(?<=[a-z])(?=[A-Z])", " ", name) or pdf_font


def restore_pdf_fonts(pdf, docx_path):
    """pdf2docx often leaves font names blank, so Word falls back to a default font.
    Put back the fonts the PDF actually used."""
    import fitz
    by_text, weight = {}, {}
    for page in fitz.open(str(pdf)):
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                for span in line["spans"]:
                    fam, text = font_family(span["font"]), span["text"].strip()
                    if text:
                        by_text.setdefault(text, fam)
                        weight[fam] = weight.get(fam, 0) + len(text)
    if not weight:
        return
    default = max(weight, key=weight.get)
    doc = Document(str(docx_path))
    for r in doc.element.body.iter(W_R):
        fonts = r.find(qn("w:rPr") + "/" + qn("w:rFonts"))
        if fonts is None or fonts.get(qn("w:ascii")):
            continue
        fam = by_text.get(run_text(r).strip(), default)
        if fam.lower().startswith(("symbol", "wingding", "zapf", "dingbat")):
            fam = default  # symbol fonts garble ordinary "•" characters in Word
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attr), fam)
    doc.save(str(docx_path))


def from_docx(src, dst):
    ext = dst.suffix.lower()
    if ext == ".docx":
        shutil.copy(src, dst)
    elif ext == ".pdf" and soffice():
        with tempfile.TemporaryDirectory() as tmp:
            run([soffice(), "--headless", "--convert-to", "pdf", "--outdir", tmp, str(src)])
            shutil.move(str(Path(tmp) / (src.stem + ".pdf")), str(dst))
    elif ext in WORD_FORMATS and WORD_APP.exists():
        word_save_as(src, dst)
    elif ext == ".pdf":
        sys.exit("PDF output needs Microsoft Word or LibreOffice installed.")
    elif ext in TEXTUTIL_EXTS:
        run(["textutil", "-convert", ext[1:], str(src), "-output", str(dst)])
    elif ext == ".pages":
        pages_import(src, dst)
    else:
        raise ValueError(ext)


def osascript(script, *args):
    run(["osascript", "-e", script, *map(str, args)], timeout=300)


def pages_convert(src, dst, save_line):
    """Drive Pages from its sandbox folder. If Pages shows a dialog (first launch,
    missing fonts), this times out; export the resume to .docx by hand instead."""
    box = Path.home() / "Library/Containers/com.apple.iWork.Pages/Data/tailor-resume"
    box.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=box) as tmp:
        tmp_src = Path(tmp) / ("in" + src.suffix)
        tmp_dst = Path(tmp) / ("out" + dst.suffix)
        shutil.copy(src, tmp_src)
        try:
            osascript(f'''on run argv
              tell application "Pages"
                open POSIX file (item 1 of argv)
                delay 2
                set d to front document
                {save_line}
                close d saving no
              end tell
            end run''', tmp_src, tmp_dst)
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as e:
            raise RuntimeError("Pages automation failed (open Pages once to clear any "
                               "dialogs, or export your resume to .docx): " + str(e))
        shutil.move(str(tmp_dst), str(dst))


def pages_export(src, dst):
    pages_convert(src, dst, "export d to POSIX file (item 2 of argv) as Microsoft Word")


def pages_import(src, dst):
    pages_convert(src, dst, "save d in POSIX file (item 2 of argv)")


# ---------- paragraph model for .docx ----------

def paragraphs(doc):
    """Every body paragraph (incl. tables and text boxes) in document order."""
    return [p for p in doc.element.body.iter(W_P)
            if not any(a.tag == MC_FALLBACK for a in p.iterancestors())]


def run_text(r):
    return "".join(TEXT_TAGS[c.tag] if TEXT_TAGS[c.tag] is not None else (c.text or "")
                   for c in r if c.tag in TEXT_TAGS)


def segments(p):
    """Group the paragraph's text runs into segments of identical formatting."""
    segs, key = [], None
    for r in p.iter(W_R):
        if next(r.iterancestors(W_P)) is not p or r.find(qn("w:drawing")) is not None:
            continue  # belongs to a nested text box, or is an image
        if not any(c.tag in TEXT_TAGS for c in r):
            continue
        rpr = r.find(qn("w:rPr"))
        k = (id(r.getparent()), rpr_key(rpr))
        if k != key:
            segs.append([])
            key = k
        segs[-1].append(r)
    return segs


def rpr_key(rpr):
    if rpr is None:
        return b""
    rpr = copy.deepcopy(rpr)
    for tag in ("w:lang", "w:noProof"):  # differences you can't see
        for el in rpr.findall(qn(tag)):
            rpr.remove(el)
    return etree.tostring(rpr)


BULLET_CHARS = set("•◦▪▫‣∙·●○■□–-*>")


def bullet_row(p):
    """If p is the text of a table row whose other cells hold only a bullet
    symbol (how pdf2docx lays out bullets), return that row."""
    cell = p.getparent()
    if cell.tag != qn("w:tc"):
        return None
    if any(q is not p and "".join(q.itertext()).strip() for q in cell.findall(W_P)):
        return None
    row = cell.getparent()
    others = ["".join(c.itertext()).strip() for c in row.findall(qn("w:tc")) if c is not cell]
    if others and all(set(t) <= BULLET_CHARS for t in others) and any(others):
        return row
    return None


def is_bullet(p):
    if bullet_row(p) is not None:
        return True
    ppr = p.find(qn("w:pPr"))
    if ppr is None:
        return False
    if ppr.find(qn("w:numPr")) is not None:
        return True
    style = ppr.find(qn("w:pStyle"))
    return style is not None and "list" in (style.get(qn("w:val")) or "").lower()


def set_run_text(r, text):
    for c in [c for c in r if c.tag in TEXT_TAGS]:
        r.remove(c)
    buf = ""
    for ch in text + "\0":
        if ch in "\t\n\0":
            if buf:
                t = r.makeelement(W_T, {})
                t.text = buf
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                r.append(t)
                buf = ""
            if ch != "\0":
                r.append(r.makeelement(qn("w:tab" if ch == "\t" else "w:br"), {}))
        else:
            buf += ch


def set_paragraph_text(p, value):
    segs = segments(p)
    if not segs:
        raise ValueError("paragraph has no editable text")
    if isinstance(value, str) or len(value) != len(segs):
        value = ["".join(value) if not isinstance(value, str) else value] + [""] * (len(segs) - 1)
    for seg, text in zip(segs, value):
        set_run_text(seg[0], text)
        for r in seg[1:]:
            r.getparent().remove(r)
        if text == "" and seg[0].getparent() is not None:
            seg[0].getparent().remove(seg[0])


def extract(docx_path):
    doc = Document(str(docx_path))
    out = []
    for i, p in enumerate(paragraphs(doc)):
        segs = segments(p)
        texts = ["".join(run_text(r) for r in s) for s in segs]
        text = "".join(texts)
        if not text.strip():
            continue
        item = {"id": i, "text": text}
        if is_bullet(p):
            item["bullet"] = True
        if len(texts) > 1:
            item["runs"] = texts
        out.append(item)
    return out


def apply_edits(docx_path, edits, out_path):
    doc = Document(str(docx_path))
    ps = paragraphs(doc)

    def get(i):
        i = int(i)
        if not 0 <= i < len(ps):
            raise ValueError(f"no paragraph with id {i}")
        return ps[i]

    for i, value in (edits.get("replace") or {}).items():
        set_paragraph_text(get(i), value)

    for group in edits.get("reorder") or []:
        els = [get(i) for i in group]
        slots = sorted(els, key=ps.index)
        marks = []
        for el in slots:
            m = el.makeelement(qn("w:bookmarkStart"), {})
            el.addprevious(m)
            marks.append(m)
        for el in els:
            el.getparent().remove(el)
        for m, el in zip(marks, els):
            m.addprevious(el)
            m.getparent().remove(m)

    for i in edits.get("delete") or []:
        p = get(i)
        parent = p.getparent()
        row = bullet_row(p)
        if row is not None:  # drop the bullet symbol along with its text
            table = row.getparent()
            table.remove(row)
            if table.find(qn("w:tr")) is None:
                table.getparent().remove(table)
        elif parent.tag == qn("w:tc") and len(parent.findall(W_P)) == 1:
            set_paragraph_text(p, "")  # a table cell must keep one paragraph
        else:
            parent.remove(p)

    doc.save(str(out_path))


# ---------- commands ----------

def prepare():
    base, ext = find_base()
    WORK.mkdir(exist_ok=True)
    info = {"base": str(base.relative_to(ROOT)), "ext": ext,
            "mode": "blocks" if ext in BLOCK_EXTS else "text"}
    if info["mode"] == "blocks":
        work_docx = WORK / "base.docx"
        to_docx(base, work_docx)
        content = extract(work_docx)
        (WORK / "content.json").write_text(json.dumps(
            {"base": info["base"], "paragraphs": content}, indent=1, ensure_ascii=False))
        info["content"] = "base/.work/content.json"
    (WORK / "info.json").write_text(json.dumps(info, indent=1))
    print(json.dumps(info))


def resume_name(d, ext):
    """carter_lee_<company>_resume<ext>, with the company taken from the agent's job.json."""
    job = d / "job.json"
    company = json.loads(job.read_text()).get("company", "") if job.exists() else ""
    company = re.sub(r"[^a-z0-9]+", "_", company.lower()).strip("_")
    if not company:
        print(f"WARN {d.relative_to(ROOT)}: no company in job.json, using the folder name")
        company = d.name.replace("-", "_")
    return f"{NAME_PREFIX}_{company}_resume{ext}"


def build(slugs):
    info_path = WORK / "info.json"
    if not info_path.exists():
        sys.exit("Run `python3 tailor.py prepare` first.")
    info = json.loads(info_path.read_text())
    ext = info["ext"]
    dirs = [OUTPUT / s for s in slugs] if slugs else sorted(
        d for d in OUTPUT.iterdir() if d.is_dir())
    for d in dirs:
        try:
            out = d / resume_name(d, ext)
            if info["mode"] == "blocks":
                edits = json.loads((d / "edits.json").read_text())
                with tempfile.TemporaryDirectory() as tmp:
                    tailored = Path(tmp) / f"{d.name}-resume.docx"
                    apply_edits(WORK / "base.docx", edits, tailored)
                    from_docx(tailored, out)
            elif (d / f"tailored{ext}").exists():
                shutil.move(str(d / f"tailored{ext}"), str(out))
            elif not out.exists():
                raise FileNotFoundError(f"agent did not write tailored{ext}")
            if ext == ".json":  # structured JSON: also render a readable .docx
                from build_resume import build as render
                render(json.loads(out.read_text()), out.with_suffix(".docx"))
            print(f"OK   {out.relative_to(ROOT)}")
        except Exception as e:
            print(f"FAIL {d.relative_to(ROOT)}: {e}")


if __name__ == "__main__":
    cmd, *rest = sys.argv[1:] or ["help"]
    if cmd == "prepare":
        prepare()
    elif cmd == "build":
        build(rest)
    else:
        print(__doc__)
